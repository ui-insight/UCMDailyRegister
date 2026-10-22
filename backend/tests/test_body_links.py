"""Tests for reading inline links and exporting them as Word hyperlinks (#408)."""

import io
from datetime import date

import pytest
from docx import Document
from docx.oxml.ns import qn
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.section import NewsletterSection
from app.utils.body_links import BodySegment, is_safe_destination, parse_body
from app.utils.export import export_newsletter_docx
from tests.conftest import make_submission_data


def docx_runs(content: bytes) -> list[tuple[str, str | None]]:
    """Every text run in document order, with its hyperlink destination if any."""
    document = Document(io.BytesIO(content))
    rels = document.part.rels
    runs: list[tuple[str, str | None]] = []
    for paragraph in document.paragraphs:
        for child in paragraph._p:
            if child.tag == qn("w:hyperlink"):
                text = "".join(t.text or "" for t in child.iter(qn("w:t")))
                runs.append((text, rels[child.get(qn("r:id"))].target_ref))
            elif child.tag == qn("w:r"):
                text = "".join(t.text or "" for t in child.iter(qn("w:t")))
                if text:
                    runs.append((text, None))
    return runs


def links_in(content: bytes) -> list[tuple[str, str]]:
    return [(text, href) for text, href in docx_runs(content) if href]


class TestParseBody:
    def test_plain_text(self):
        assert parse_body("No links here.") == [BodySegment("No links here.")]

    def test_link_keeps_its_exact_words_and_position(self):
        assert parse_body('Learn more today. <a href="https://uidaho.edu/x">Learn more</a>.') == [
            BodySegment("Learn more today. "),
            BodySegment("Learn more", "https://uidaho.edu/x"),
            BodySegment("."),
        ]

    @pytest.mark.parametrize(
        "anchor",
        [
            "<a href='https://uidaho.edu'>site</a>",
            '<a target="_blank" href="https://uidaho.edu" rel="noopener">site</a>',
            '<A HREF="https://uidaho.edu">site</A>',
            "<a href=https://uidaho.edu>site</a>",
        ],
    )
    def test_tolerates_attribute_and_quote_variations(self, anchor: str):
        assert parse_body(anchor) == [BodySegment("site", "https://uidaho.edu")]

    def test_destination_entities_are_decoded(self):
        (segment,) = parse_body('<a href="https://uidaho.edu/?a=1&amp;b=2">form</a>')
        assert segment.href == "https://uidaho.edu/?a=1&b=2"

    def test_mailto_link(self):
        (segment,) = parse_body('<a href="mailto:ucm@uidaho.edu">UCM</a>')
        assert segment == BodySegment("UCM", "mailto:ucm@uidaho.edu")

    @pytest.mark.parametrize(
        "href", ["javascript:alert(1)", "http://uidaho.edu", "https://user:pw@uidaho.edu", "ftp://x"]
    )
    def test_unsafe_destination_degrades_to_plain_words(self, href: str):
        assert parse_body(f'Go <a href="{href}">here</a>') == [
            BodySegment("Go "),
            BodySegment("here"),
        ]

    def test_stray_anchor_tags_never_leak(self):
        segments = parse_body('Broken <a href="https://uidaho.edu">link and a < b')
        assert "".join(s.text for s in segments) == "Broken link and a < b"
        assert all(s.href is None for s in segments)


class TestSafeDestination:
    @pytest.mark.parametrize(
        ("url", "safe"),
        [
            ("https://www.uidaho.edu/events", True),
            ("mailto:jbauer@uidaho.edu", True),
            ("mailto:jbauer@uidaho.edu?subject=hi", False),
            ("https://", False),
            ("uidaho.edu", False),
        ],
    )
    def test_rules_match_the_editor(self, url: str, safe: bool):
        assert is_safe_destination(url) is safe


class TestDocxHyperlinks:
    def _export(self, body: str, headline: str = "Headline") -> bytes:
        buffer = export_newsletter_docx(
            newsletter_type="tdr",
            publish_date=date(2026, 11, 2),
            sections=[{"name": "News", "items": [{"Final_Headline": headline, "Final_Body": body}]}],
        )
        return buffer.getvalue()

    def test_link_is_a_real_hyperlink_with_its_destination(self):
        content = self._export('Register at <a href="https://uidaho.edu/register">the form</a>.')
        assert links_in(content) == [("the form", "https://uidaho.edu/register")]

    def test_link_stays_on_the_chosen_occurrence(self):
        content = self._export(
            'Learn more about the fair. Questions? <a href="https://uidaho.edu/fair">Learn more</a>.'
        )
        runs = docx_runs(content)
        start = next(i for i, (text, _) in enumerate(runs) if text.startswith("Learn more about"))
        assert runs[start:start + 3] == [
            ("Learn more about the fair. Questions? ", None),
            ("Learn more", "https://uidaho.edu/fair"),
            (".", None),
        ]

    def test_unsafe_link_exports_as_plain_words_without_html(self):
        content = self._export('Click <a href="javascript:alert(1)">here</a>.')
        assert links_in(content) == []
        text = "".join(t for t, _ in docx_runs(content))
        assert "Click here." in text
        assert "<a" not in text

    def test_multiline_link_text_keeps_its_line_break(self):
        content = self._export('<a href="https://uidaho.edu">Apply\nnow</a>')
        document = Document(io.BytesIO(content))
        (hyperlink,) = document.element.body.iter(qn("w:hyperlink"))
        assert len(list(hyperlink.iter(qn("w:br")))) == 1


@pytest.mark.asyncio
async def test_approved_link_reaches_the_exported_word_file(
    client: AsyncClient, db: AsyncSession, staff_headers: dict[str, str]
):
    db.add(
        NewsletterSection(
            Newsletter_Type="tdr",
            Name="Employee News",
            Slug="employee-news",
            Display_Order=1,
            Is_Active=True,
        )
    )
    await db.commit()

    resp = await client.post(
        "/api/v1/submissions/",
        json=make_submission_data(Schedule_Requests=[{"Requested_Date": "2026-11-02"}]),
        headers=staff_headers,
    )
    submission_id = resp.json()["Id"]
    body = (
        "Email the Library about extended hours. "
        'Questions? <a href="mailto:library@uidaho.edu">Email the Library</a> or see '
        '<a href="https://www.lib.uidaho.edu/hours">the hours page</a>.'
    )
    resp = await client.post(
        f"/api/v1/ai-edits/{submission_id}/finalize",
        json={"Headline": "Library extends hours", "Body": body, "Approve_For_Newsletter": True},
        headers=staff_headers,
    )
    assert resp.status_code == 200

    resp = await client.post(
        "/api/v1/newsletters/assemble",
        json={"Newsletter_Type": "tdr", "Publish_Date": "2026-11-02"},
        headers=staff_headers,
    )
    newsletter_id = resp.json()["Id"]
    resp = await client.get(f"/api/v1/newsletters/{newsletter_id}/export", headers=staff_headers)
    assert resp.status_code == 200

    runs = docx_runs(resp.content)
    assert links_in(resp.content) == [
        ("Email the Library", "mailto:library@uidaho.edu"),
        ("the hours page", "https://www.lib.uidaho.edu/hours"),
    ]
    # The earlier, informational mention stays plain text.
    assert ("Email the Library about extended hours. Questions? ", None) in runs
